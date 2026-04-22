from __future__ import annotations

import importlib.util
import io
import re
from copy import deepcopy
from pathlib import Path
from typing import Iterable

from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from src.adapters.pdf_adapter import PDFAdapter


MISSING_PILLOW_MSG = (
    "Para extraer imágenes es necesario instalar la librería Pillow. Ejecuta: pip install pillow"
)


class PDFService:
<<<<<<< HEAD
    """Casos de uso de negocio de procesamiento PDF."""

    def __init__(self, pdf_adapter: PDFAdapter | None = None) -> None:
        self._pdf_adapter = pdf_adapter or PDFAdapter()

    def _image_extraction_available(self) -> bool:
        return importlib.util.find_spec("PIL") is not None

    def _ensure_image_support(self) -> None:
=======
    """
    Casos de uso de negocio de procesamiento PDF.
    Contiene la lógica para unir, dividir, extraer texto y convertir PDFs a Word.
    """

    def __init__(self, pdf_adapter: PDFAdapter | None = None) -> None:
        """
        Inicializa el servicio con un adaptador de PDF.
        :param pdf_adapter: Adaptador que implementa las operaciones de bajo nivel (pypdf).
        """
        self._pdf_adapter = pdf_adapter or PDFAdapter()

    def _image_extraction_available(self) -> bool:
        """Verifica si la librería Pillow está disponible para extracción de imágenes."""
        return importlib.util.find_spec("PIL") is not None

    def _ensure_image_support(self) -> None:
        """Lanza un error si no hay soporte para imágenes."""
>>>>>>> main
        if not self._image_extraction_available():
            raise RuntimeError(MISSING_PILLOW_MSG)

    def get_total_pages(self, input_path: Path) -> int:
<<<<<<< HEAD
=======
        """
        Obtiene el número total de páginas de un PDF.
        :param input_path: Ruta al archivo PDF.
        :return: Cantidad de páginas.
        """
>>>>>>> main
        reader = self._pdf_adapter.reader(input_path)
        return len(reader.pages)

    def merge_pdfs(self, input_paths: Iterable[Path], output_path: Path) -> Path:
<<<<<<< HEAD
=======
        """
        Une múltiples archivos PDF en uno solo respetando el orden.
        :param input_paths: Lista de rutas de los PDFs a unir.
        :param output_path: Ruta donde se guardará el PDF resultante.
        :return: Ruta del archivo generado.
        """
>>>>>>> main
        writer = self._pdf_adapter.writer()
        for path in input_paths:
            reader = self._pdf_adapter.reader(path)
            for page in reader.pages:
                writer.add_page(page)
        with output_path.open("wb") as f:
            writer.write(f)
        return output_path

    def split_pdf(
        self,
        input_path: Path,
        output_dir: Path,
        start_page: int | None = None,
        end_page: int | None = None,
    ) -> list[Path]:
<<<<<<< HEAD
        reader = self._pdf_adapter.reader(input_path)
        total_pages = len(reader.pages)
        if total_pages == 0:
            raise ValueError("El PDF no contiene páginas")

        start = 1 if start_page is None else start_page
        end = total_pages if end_page is None else end_page

        if start < 1 or end > total_pages or start > end:
            raise ValueError(
                f"Rango inválido. Debe estar entre 1 y {total_pages}, recibido {start}-{end}"
            )
=======
        """
        Divide un PDF por un rango de páginas específico.
        Ajusta automáticamente el rango si excede los límites del documento (clamping).
        :param input_path: PDF de origen.
        :param output_dir: Carpeta de destino.
        :param start_page: Página inicial (1-based).
        :param end_page: Página final.
        :return: Lista de rutas de los archivos generados.
        """
        reader = self._pdf_adapter.reader(input_path)
        total_pages = len(reader.pages)
        if total_pages == 0:
            return []

        # Ajuste flexible de límites (clamping)
        start = max(1, start_page if start_page is not None else 1)
        end = min(total_pages, end_page if end_page is not None else total_pages)

        # Si el inicio está fuera del documento, no se puede dividir
        if start > total_pages:
            return []
>>>>>>> main

        output_dir.mkdir(parents=True, exist_ok=True)
        generated: list[Path] = []

        for page_number in range(start, end + 1):
            writer = self._pdf_adapter.writer()
            writer.add_page(reader.pages[page_number - 1])
            out_file = output_dir / f"{input_path.stem}_pagina_{page_number}.pdf"
            with out_file.open("wb") as f:
                writer.write(f)
            generated.append(out_file)

        return generated

    def split_pdf_by_parts(self, input_path: Path, output_dir: Path, num_parts: int) -> list[Path]:
<<<<<<< HEAD
        reader = self._pdf_adapter.reader(input_path)
        total_pages = len(reader.pages)
        if total_pages == 0:
            raise ValueError("El PDF no contiene páginas")
        if num_parts < 1 or num_parts > total_pages:
            raise ValueError(
                f"Número de partes inválido. Debe estar entre 1 y {total_pages}, recibido {num_parts}"
            )

        base_pages_per_part = total_pages // num_parts
        remainder = total_pages % num_parts
=======
        """
        Divide un PDF en un número determinado de partes equitativas.
        Si el PDF tiene menos páginas que partes, se divide en el máximo posible (1 página por parte).
        :param input_path: PDF de origen.
        :param output_dir: Carpeta de destino.
        :param num_parts: Número de partes deseadas.
        :return: Lista de rutas de los archivos generados.
        """
        reader = self._pdf_adapter.reader(input_path)
        total_pages = len(reader.pages)
        if total_pages == 0:
            return []

        # Ajuste flexible de partes
        actual_parts = min(num_parts, total_pages)
        if actual_parts < 1:
            return []

        base_pages_per_part = total_pages // actual_parts
        remainder = total_pages % actual_parts
>>>>>>> main

        output_dir.mkdir(parents=True, exist_ok=True)
        generated: list[Path] = []
        current_page = 0

<<<<<<< HEAD
        for part_index in range(1, num_parts + 1):
=======
        for part_index in range(1, actual_parts + 1):
>>>>>>> main
            pages_in_this_part = base_pages_per_part + (1 if part_index <= remainder else 0)
            writer = self._pdf_adapter.writer()

            for _ in range(pages_in_this_part):
                writer.add_page(reader.pages[current_page])
                current_page += 1

            start_page = current_page - pages_in_this_part + 1
            end_page = current_page
            out_file = output_dir / f"{input_path.stem}_parte_{part_index}_{start_page}-{end_page}.pdf"
            with out_file.open("wb") as f:
                writer.write(f)
            generated.append(out_file)

        return generated

<<<<<<< HEAD
    def parse_page_ranges(self, ranges_input: str, total_pages: int) -> list[tuple[int, int]]:
=======
    def parse_page_ranges(self, ranges_input: str, total_pages: int, clamp: bool = False) -> list[tuple[int, int]]:
        """
        Parsea una cadena de rangos (ej: '1-5, 8, 10-12').
        :param ranges_input: Cadena de entrada del usuario.
        :param total_pages: Total de páginas para validación.
        :param clamp: Si es True, ajusta los rangos a los límites del PDF en lugar de lanzar error.
        :return: Lista de tuplas (inicio, fin).
        """
>>>>>>> main
        if not ranges_input.strip():
            raise ValueError("Debes ingresar al menos un rango.")

        parsed_ranges: list[tuple[int, int]] = []

        for chunk in ranges_input.split(","):
            raw = chunk.strip()
            if not raw:
                continue

            if "-" in raw:
                pieces = raw.split("-", maxsplit=1)
                start = int(pieces[0].strip())
                end = int(pieces[1].strip())
            else:
                start = int(raw)
                end = start

<<<<<<< HEAD
            if start < 1 or end > total_pages:
                raise ValueError(f"Rango fuera de límites: {start}-{end}. Total de páginas: {total_pages}")
            if start > end:
                raise ValueError(f"Rango inválido: {start}-{end} (inicio mayor que fin)")

            parsed_ranges.append((start, end))

        if not parsed_ranges:
            raise ValueError("No se detectaron rangos válidos.")

        occupied_pages: set[int] = set()
        for start, end in parsed_ranges:
            page_set = set(range(start, end + 1))
            if occupied_pages.intersection(page_set):
                raise ValueError("No se permiten rangos solapados o repetidos.")
            occupied_pages.update(page_set)

        return parsed_ranges

    def extract_page_ranges(self, input_path: Path, output_dir: Path, ranges_input: str) -> list[Path]:
        reader = self._pdf_adapter.reader(input_path)
        total_pages = len(reader.pages)
        if total_pages == 0:
            raise ValueError("El PDF no contiene páginas")

        ranges = self.parse_page_ranges(ranges_input=ranges_input, total_pages=total_pages)
=======
            if clamp:
                # Ajustar a límites del documento
                if start > total_pages:
                    continue # Rango completamente fuera
                start = max(1, start)
                end = min(total_pages, end)
                if start > end:
                    continue
            else:
                if start < 1 or end > total_pages:
                    raise ValueError(f"Rango fuera de límites: {start}-{end}. Total de páginas: {total_pages}")
                if start > end:
                    raise ValueError(f"Rango inválido: {start}-{end} (inicio mayor que fin)")

            parsed_ranges.append((start, end))

        if not parsed_ranges and not clamp:
            raise ValueError("No se detectaron rangos válidos.")

        # Eliminamos validación de solapamiento estricta para permitir mayor flexibilidad en batch
        return parsed_ranges

    def extract_page_ranges(self, input_path: Path, output_dir: Path, ranges_input: str) -> list[Path]:
        """
        Extrae múltiples rangos de páginas en archivos independientes.
        Ajusta los rangos dinámicamente según la longitud de cada PDF.
        :param input_path: PDF de origen.
        :param output_dir: Carpeta de destino.
        :param ranges_input: Cadena con los rangos (ej: '2-14, 16-18').
        :return: Lista de archivos generados.
        """
        reader = self._pdf_adapter.reader(input_path)
        total_pages = len(reader.pages)
        if total_pages == 0:
            return []

        # Usamos clamp=True para que sea flexible con archivos de distinta longitud
        ranges = self.parse_page_ranges(ranges_input=ranges_input, total_pages=total_pages, clamp=True)
>>>>>>> main
        output_dir.mkdir(parents=True, exist_ok=True)

        generated: list[Path] = []
        for index, (start, end) in enumerate(ranges, start=1):
            writer = self._pdf_adapter.writer()
            for page_number in range(start, end + 1):
                writer.add_page(reader.pages[page_number - 1])

            out_file = output_dir / f"{input_path.stem}_rango_{index}_{start}-{end}.pdf"
            with out_file.open("wb") as f:
                writer.write(f)
            generated.append(out_file)

        return generated

    def extract_text(self, input_path: Path) -> str:
<<<<<<< HEAD
=======
        """Extrae todo el texto plano de un PDF."""
>>>>>>> main
        reader = self._pdf_adapter.reader(input_path)
        text_parts: list[str] = []
        for page in reader.pages:
            text_parts.append(page.extract_text() or "")
        return "\n".join(text_parts).strip()

    def extract_text_and_images(self, input_path: Path, output_dir: Path) -> dict[str, Path | int]:
<<<<<<< HEAD
=======
        """
        Extrae texto a un archivo .txt e imágenes a una carpeta 'images'.
        :param input_path: PDF de origen.
        :param output_dir: Carpeta de destino del contenido.
        :return: Diccionario con rutas y conteo de imágenes.
        """
>>>>>>> main
        self._ensure_image_support()

        reader = self._pdf_adapter.reader(input_path)
        output_dir.mkdir(parents=True, exist_ok=True)
        images_dir = output_dir / "images"
        images_dir.mkdir(parents=True, exist_ok=True)

        text = self.extract_text(input_path)
        text_file = output_dir / f"{input_path.stem}_texto.txt"
        text_file.write_text(text, encoding="utf-8")

        image_count = 0
        for page_index, page in enumerate(reader.pages, start=1):
            try:
                page_images = page.images
            except ImportError as exc:
                raise RuntimeError(MISSING_PILLOW_MSG) from exc

            for image_index, image_file in enumerate(page_images, start=1):
                suffix = Path(image_file.name).suffix or ".png"
                image_path = images_dir / f"pagina_{page_index:03d}_{image_index:03d}{suffix}"
                image_path.write_bytes(image_file.data)
                image_count += 1

        return {
            "text_file": text_file,
            "images_dir": images_dir,
            "image_count": image_count,
        }

    def _add_hyperlink(self, paragraph, url: str, text: str) -> None:
<<<<<<< HEAD
=======
        """Añade un hipervínculo funcional a un párrafo de python-docx."""
>>>>>>> main
        part = paragraph.part
        r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", True)

        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("r:id"), r_id)
        new_run = OxmlElement("w:r")
        r_pr = OxmlElement("w:rPr")

        color = OxmlElement("w:color")
        color.set(qn("w:val"), "0000FF")
        r_pr.append(color)
        underline = OxmlElement("w:u")
        underline.set(qn("w:val"), "single")
        r_pr.append(underline)

        new_run.append(r_pr)
        w_text = OxmlElement("w:t")
        w_text.text = text
        new_run.append(w_text)
        hyperlink.append(new_run)
        paragraph._p.append(hyperlink)

    def _write_text_with_links(self, paragraph, text: str) -> None:
<<<<<<< HEAD
=======
        """Detecta URLs en texto y las escribe como links en el párrafo."""
>>>>>>> main
        url_pattern = re.compile(r"https?://\S+")
        cursor = 0
        for match in url_pattern.finditer(text):
            if match.start() > cursor:
                paragraph.add_run(text[cursor : match.start()])
            url = match.group(0)
            self._add_hyperlink(paragraph, url, url)
            cursor = match.end()
        if cursor < len(text):
            paragraph.add_run(text[cursor:])

    def _convert_pdf_to_docx_quick(self, input_path: Path, output_docx: Path) -> Path:
<<<<<<< HEAD
=======
        """Conversión básica a Word extrayendo texto e imágenes secuencialmente."""
>>>>>>> main
        from docx import Document

        reader = self._pdf_adapter.reader(input_path)
        document = Document()

        has_any_text = False
        image_support = self._image_extraction_available()

        for page_index, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            if text:
                has_any_text = True
                document.add_heading(f"Página {page_index}", level=2)
                for paragraph in text.split("\n"):
                    p = document.add_paragraph()
                    self._write_text_with_links(p, paragraph)

            if image_support:
                try:
                    page_images = page.images
                except ImportError:
                    page_images = []
                    image_support = False
                for image_file in page_images:
                    image_stream = io.BytesIO(image_file.data)
                    try:
                        document.add_picture(image_stream)
                    except Exception:
                        continue

        if not has_any_text:
            document.add_paragraph(
                "[No se encontró texto extraíble en el PDF. Puede ser un PDF escaneado (sin OCR).]"
            )
        if not image_support:
            document.add_paragraph(f"[Imágenes omitidas: {MISSING_PILLOW_MSG}]")

        output_docx.parent.mkdir(parents=True, exist_ok=True)
        document.save(str(output_docx))
        return output_docx

    def _convert_pdf_to_docx_structured_fallback(self, input_path: Path, output_docx: Path) -> Path:
<<<<<<< HEAD
=======
        """Conversión estructurada usando PyMuPDF para detectar bloques y pdfplumber para tablas."""
>>>>>>> main
        from docx import Document

        try:
            import fitz
        except ImportError:
            return self._convert_pdf_to_docx_quick(input_path, output_docx)

        document = Document()
        pdf = fitz.open(str(input_path))

        plumber_doc = None
        try:
            import pdfplumber

            plumber_doc = pdfplumber.open(str(input_path))
        except Exception:
            plumber_doc = None

        for page_index, page in enumerate(pdf, start=1):
            document.add_heading(f"Página {page_index}", level=1)
            blocks: list[tuple[float, float, str, object]] = []

            # Tables via pdfplumber
            if plumber_doc is not None:
                try:
                    plumber_page = plumber_doc.pages[page_index - 1]
                    for table in plumber_page.find_tables():
                        blocks.append((table.bbox[1], table.bbox[0], "table", table.extract()))
                except Exception:
                    pass

            page_dict = page.get_text("dict")
            for block in page_dict.get("blocks", []):
                if block.get("type") == 0:  # text
                    text_lines: list[str] = []
                    max_size = 0.0
                    is_bold = False
                    is_italic = False
                    for line in block.get("lines", []):
                        line_text = ""
                        for span in line.get("spans", []):
                            span_text = span.get("text", "")
                            line_text += span_text
                            max_size = max(max_size, float(span.get("size", 0)))
                            flags = int(span.get("flags", 0))
                            is_italic = is_italic or bool(flags & 2)
                            is_bold = is_bold or bool(flags & 16)
                        if line_text.strip():
                            text_lines.append(line_text)
                    text_content = "\n".join(text_lines).strip()
                    if text_content:
                        blocks.append((block["bbox"][1], block["bbox"][0], "text", (text_content, max_size, is_bold, is_italic)))
                elif block.get("type") == 1:  # image
                    blocks.append((block["bbox"][1], block["bbox"][0], "image", block))

            blocks.sort(key=lambda x: (x[0], x[1]))

            for _, _, block_type, payload in blocks:
                if block_type == "text":
                    text_content, max_size, is_bold, is_italic = payload
                    level = 0 if max_size >= 16 else (1 if max_size >= 13 else None)
                    paragraph = document.add_heading(level=level) if level is not None else document.add_paragraph()
                    self._write_text_with_links(paragraph, text_content)
                    for run in paragraph.runs:
                        run.bold = is_bold
                        run.italic = is_italic
                elif block_type == "table":
                    table_data = payload
                    if not table_data:
                        continue
                    rows = len(table_data)
                    cols = max(len(row) if row else 0 for row in table_data)
                    if rows == 0 or cols == 0:
                        continue
                    table = document.add_table(rows=rows, cols=cols)
                    for r, row in enumerate(table_data):
                        for c in range(cols):
                            table.cell(r, c).text = (row[c] if row and c < len(row) and row[c] else "")
                else:
                    try:
                        image_bytes = page.get_image_info(xrefs=True)
                        # fallback simple: extract first image xref in page
                        image_list = page.get_images(full=True)
                        if not image_list:
                            continue
                        xref = image_list[0][0]
                        image = pdf.extract_image(xref)
                        document.add_picture(io.BytesIO(image["image"]))
                    except Exception:
                        continue

        if plumber_doc is not None:
            plumber_doc.close()
        pdf.close()

        output_docx.parent.mkdir(parents=True, exist_ok=True)
        document.save(str(output_docx))
        return output_docx

    def _insert_paragraph_after(self, paragraph):
<<<<<<< HEAD
=======
        """Inserta un nuevo párrafo justo después de uno existente."""
>>>>>>> main
        new_p = OxmlElement("w:p")
        paragraph._p.addnext(new_p)
        from docx.text.paragraph import Paragraph

        return Paragraph(new_p, paragraph._parent)

    def _fix_images_layout(self, docx_path: Path, max_width_inches: float = 6.0) -> None:
<<<<<<< HEAD
        from docx import Document

        doc = Document(str(docx_path))
        paragraphs = list(doc.paragraphs)

        for paragraph in paragraphs:
            image_runs = [run for run in paragraph.runs if run.element.xpath('.//pic:pic')]
            if not image_runs:
                continue

            has_text = bool(paragraph.text.strip())
            if has_text or len(paragraph.runs) > len(image_runs):
                current = paragraph
                for image_run in image_runs:
                    image_para = self._insert_paragraph_after(current)
                    image_para.add_run()._r.append(deepcopy(image_run._r))
                    image_run._r.getparent().remove(image_run._r)
                    current = image_para

        for shape in doc.inline_shapes:
            if shape.width > Inches(max_width_inches):
                ratio = shape.height / shape.width
                shape.width = Inches(max_width_inches)
                shape.height = int(shape.width * ratio)

        for paragraph in list(doc.paragraphs):
            has_image = any(run.element.xpath('.//pic:pic') for run in paragraph.runs)
            if has_image:
                paragraph.insert_paragraph_before()
                self._insert_paragraph_after(paragraph)
=======
        """
        Corrige problemas de diseño en Word, especialmente solapamiento de imágenes.
        Convierte imágenes flotantes a in-line y asegura flujo vertical.
        """
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
        from docx.oxml.ns import qn

        doc = Document(str(docx_path))
        
        # 1. Convertir imágenes flotantes (wp:anchor) a imágenes en línea (wp:inline)
        for drawing in doc.element.xpath('//w:drawing'):
            anchor = drawing.find(qn('wp:anchor'))
            if anchor is not None:
                inline = OxmlElement('wp:inline')
                for attr in ['distT', 'distB', 'distL', 'distR']:
                    val = anchor.get(attr)
                    if val: inline.set(attr, val)
                
                for child in anchor.getchildren():
                    tag_local = child.tag.split('}')[-1] if '}' in child.tag else child.tag
                    if tag_local not in ['simplePos', 'positionH', 'positionV', 'wrapNone', 'wrapSquare', 'wrapTight', 'wrapThrough', 'wrapTopAndBottom']:
                        inline.append(deepcopy(child))
                
                drawing.replace(anchor, inline)

        # Función auxiliar para separar imágenes de texto
        def process_paragraphs(paragraphs_list):
            for paragraph in list(paragraphs_list):
                has_image = bool(paragraph._p.xpath('.//w:drawing'))
                if not has_image:
                    continue

                if paragraph.text.strip():
                    image_runs = [run for run in paragraph.runs if run._r.xpath('.//w:drawing')]
                    if image_runs:
                        current = paragraph
                        for image_run in image_runs:
                            image_para = self._insert_paragraph_after(current)
                            image_para.add_run()._r.append(deepcopy(image_run._r))
                            image_run._r.getparent().remove(image_run._r)
                            current = image_para

        # Procesar todos los párrafos (principales y tablas)
        all_paras = list(doc.paragraphs)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    all_paras.extend(cell.paragraphs)
        
        process_paragraphs(all_paras)

        # Aplicar formato de alineación e interlineado
        all_paras_updated = list(doc.paragraphs)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    all_paras_updated.extend(cell.paragraphs)

        for paragraph in all_paras_updated:
            if paragraph._p.xpath('.//w:drawing'):
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.space_before = Inches(0.15)
                paragraph.paragraph_format.space_after = Inches(0.15)
                paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
                paragraph.paragraph_format.line_spacing = 1.0

        # Ajustar tamaño
        for shape in doc.inline_shapes:
            try:
                if shape.width > Inches(max_width_inches):
                    ratio = shape.height / shape.width
                    shape.width = Inches(max_width_inches)
                    shape.height = int(shape.width * ratio)
            except Exception:
                continue
>>>>>>> main

        doc.save(str(docx_path))

    def _convert_pdf_to_docx_pdf2docx(self, input_path: Path, output_docx: Path) -> Path:
<<<<<<< HEAD
=======
        """Motor principal de conversión usando pdf2docx."""
>>>>>>> main
        from pdf2docx import Converter

        output_docx.parent.mkdir(parents=True, exist_ok=True)
        cv = Converter(str(input_path))
        try:
            cv.convert(str(output_docx), start=0, end=None)
        finally:
            cv.close()
<<<<<<< HEAD
        self._fix_images_layout(output_docx)
        return output_docx

    def is_text_based_pdf(self, input_path: Path) -> bool:
=======
        return output_docx

    def is_text_based_pdf(self, input_path: Path) -> bool:
        """Verifica si un PDF tiene capa de texto extraíble."""
>>>>>>> main
        reader = self._pdf_adapter.reader(input_path)
        for page in reader.pages:
            if (page.extract_text() or "").strip():
                return True
        return False

    def convert_pdf_to_docx(self, input_path: Path, output_docx: Path, mode: str = "advanced") -> Path:
<<<<<<< HEAD
        if mode == "quick":
            return self._convert_pdf_to_docx_quick(input_path, output_docx)

        try:
            return self._convert_pdf_to_docx_pdf2docx(input_path, output_docx)
        except Exception:
            return self._convert_pdf_to_docx_structured_fallback(input_path, output_docx)
=======
        """
        Orquesta la conversión de PDF a Word eligiendo el motor adecuado.
        :param input_path: PDF de origen.
        :param output_docx: Ruta de salida .docx.
        :param mode: 'quick' o 'advanced'.
        :return: Ruta del archivo generado.
        """
        if mode == "quick":
            res = self._convert_pdf_to_docx_quick(input_path, output_docx)
        else:
            try:
                res = self._convert_pdf_to_docx_pdf2docx(input_path, output_docx)
            except Exception:
                res = self._convert_pdf_to_docx_structured_fallback(input_path, output_docx)
        
        self._fix_images_layout(output_docx)
        return res

    def convert_docx_to_pdf(self, input_path: Path, output_pdf: Path) -> Path:
        """
        Convierte un archivo Word (.docx) a PDF.
        :param input_path: Archivo Word de origen.
        :param output_pdf: Ruta de salida .pdf.
        :return: Ruta del archivo generado.
        """
        from docx2pdf import convert

        output_pdf.parent.mkdir(parents=True, exist_ok=True)
        # docx2pdf puede ser ruidoso, pero es efectivo en Windows con Word
        convert(str(input_path), str(output_pdf))
        return output_pdf
>>>>>>> main
