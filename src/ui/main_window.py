from __future__ import annotations

import base64
import shutil
import subprocess
import sys
import tempfile
import tkinter as tk
from pathlib import Path
from tkinter import filedialog, messagebox, scrolledtext, ttk

from docx import Document
from src.services.file_service import FileService
from src.services.pdf_service import PDFService
from src.services.viewer_service import ViewerService
from src.ui.components import create_button
from src.utils.helpers import ensure_pdf_extension, human_error


class MainWindow(ttk.Frame):
    """
    Clase principal de la interfaz gráfica.
    Gestiona la interacción del usuario con los servicios de PDF, archivos y visor.
    """

    def __init__(
        self,
        master: tk.Tk,
        file_service: FileService,
        pdf_service: PDFService,
        viewer_service: ViewerService,
    ) -> None:
        """
        Inicializa la ventana principal y configura los servicios.
        :param master: Ventana raíz de Tkinter.
        :param file_service: Servicio de archivos.
        :param pdf_service: Servicio de procesamiento PDF.
        :param viewer_service: Servicio de visualización.
        """
        super().__init__(master, padding=0)
        self.master = master
        self.file_service = file_service
        self.pdf_service = pdf_service
        self.viewer_service = viewer_service
        self.loaded_files: list[Path] = []
        self.status_var = tk.StringVar(value="Listo para procesar PDFs")
        self.pdf_info_var = tk.StringVar(value="Selecciona un PDF de la lista")
        self.preview_current_pdf: Path | None = None
        self.preview_images: list[tk.PhotoImage] = []
        self.preview_tab_display_name = "👁️ Vista"
        self.preview_temp_dir: Path | None = None
        self._spire_notice_shown = False

        self._configure_styles()
        self._build_ui()
        self._enable_fullscreen_on_start()

    def _enable_fullscreen_on_start(self) -> None:
        try:
            self.master.state("zoomed")
        except Exception:
            self.master.attributes("-zoomed", True)

    def _configure_styles(self) -> None:
        """Configura los estilos de TTK para un aspecto moderno."""
        style = ttk.Style()
        style.theme_use("clam")  # Un tema base más moderno y personalizable que el default de Windows

        # Colores
        bg_main = "#f0f2f5"
        bg_sidebar = "#ffffff"
        primary_color = "#1a73e8"
        accent_color = "#34a853"
        text_color = "#202124"

        self.master.configure(bg=bg_main)

        # Estilo de botones
        style.configure("TButton", padding=6, font=("Segoe UI", 10))
        style.configure("Primary.TButton", foreground="white", background=primary_color)
        style.map("Primary.TButton", background=[("active", "#1557b0")])
        
        style.configure("Accent.TButton", foreground="white", background=accent_color)
        style.map("Accent.TButton", background=[("active", "#2d8e47")])

        # Etiquetas
        style.configure("TLabel", background=bg_main, foreground=text_color, font=("Segoe UI", 10))
        style.configure("Header.TLabel", font=("Segoe UI", 14, "bold"), foreground=primary_color)
        style.configure("Sidebar.TFrame", background=bg_sidebar)
        style.configure("Sidebar.TLabel", background=bg_sidebar, font=("Segoe UI", 10, "bold"))

    def _build_ui(self) -> None:
        """Construye la disposición de los elementos en la ventana."""
        self.pack(fill=tk.BOTH, expand=True)

        # 1. Cabecera
        header = ttk.Frame(self, padding=(20, 10))
        header.pack(fill=tk.X)
        ttk.Label(header, text="📄 PDF Master Pro", style="Header.TLabel").pack(side=tk.LEFT)
        
        # 2. Área principal con PanedWindow
        paned = ttk.PanedWindow(self, orient=tk.HORIZONTAL)
        paned.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)

        # 2a. Sidebar (Gestión de archivos)
        sidebar = ttk.Frame(paned, padding=10, style="Sidebar.TFrame")
        paned.add(sidebar, weight=1)

        ttk.Label(sidebar, text="📂 Archivos Cargados", style="Sidebar.TLabel").pack(anchor=tk.W, pady=(0, 5))
        
        list_container = ttk.Frame(sidebar, style="Sidebar.TFrame")
        list_container.pack(fill=tk.BOTH, expand=True)

        self.files_list = tk.Listbox(
            list_container, 
            height=15, 
            font=("Segoe UI", 9), 
            borderwidth=1, 
            relief=tk.FLAT,
            selectbackground="#e8f0fe",
            selectforeground="#1a73e8"
        )
        self.files_list.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        self.files_list.bind("<<ListboxSelect>>", self._on_file_selection)

        scrollbar = ttk.Scrollbar(list_container, orient=tk.VERTICAL, command=self.files_list.yview)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        self.files_list.config(yscrollcommand=scrollbar.set)

        # Controles de gestión de archivos (Subir, Bajar, Quitar, Limpiar)
        file_mgmt_frame = ttk.Frame(sidebar, style="Sidebar.TFrame", padding=(0, 5))
        file_mgmt_frame.pack(fill=tk.X)
        file_mgmt_frame.columnconfigure((0, 1), weight=1)

        create_button(file_mgmt_frame, "🔼 Subir", self._move_selected_up).grid(row=0, column=0, sticky="ew", padx=(0, 2), pady=2)
        create_button(file_mgmt_frame, "🔽 Bajar", self._move_selected_down).grid(row=0, column=1, sticky="ew", padx=(2, 0), pady=2)
        create_button(file_mgmt_frame, "❌ Quitar", self._remove_selected_file).grid(row=1, column=0, sticky="ew", padx=(0, 2), pady=2)
        create_button(file_mgmt_frame, "🧹 Limpiar", self._clear_file_list).grid(row=1, column=1, sticky="ew", padx=(2, 0), pady=2)

        # 2b. Área de contenido (Acciones y Consola)
        content_area = ttk.Frame(paned, padding=(15, 0))
        paned.add(content_area, weight=3)

        # Acciones Principales (Organizadas por grupos)
        actions_frame = ttk.LabelFrame(content_area, text="🛠️ Herramientas", padding=15)
        actions_frame.pack(fill=tk.X, pady=(0, 15))

        # Fila 1: Gestión y Operaciones PDF
        row1 = ttk.Frame(actions_frame)
        row1.pack(fill=tk.X, pady=(0, 5))
        row1.columnconfigure((0, 1, 2), weight=1)
        
        create_button(row1, "📥 Cargar Archivos", self._load_files, style="Primary.TButton").grid(row=0, column=0, sticky="ew", padx=5)
        create_button(row1, "🔗 Unir", self._merge_pdfs, style="Accent.TButton").grid(row=0, column=1, sticky="ew", padx=5)
        create_button(row1, "✂️ Dividir", self._split_pdf).grid(row=0, column=2, sticky="ew", padx=5)

        # Fila 2: Conversiones
        row2 = ttk.Frame(actions_frame)
        row2.pack(fill=tk.X, pady=(5, 0))
        row2.columnconfigure((0, 1, 2), weight=1)
        
        create_button(row2, "📝 PDF a Word / Extraer", self._extract_content).grid(row=0, column=0, sticky="ew", padx=5)
        create_button(row2, "📘 Word a PDF", self._word_to_pdf).grid(row=0, column=1, sticky="ew", padx=5)
        create_button(row2, "✍️ Editar / Firmar PDF", self._edit_pdf).grid(row=0, column=2, sticky="ew", padx=5)

        # Información del PDF seleccionado
        info_frame = ttk.Frame(content_area, padding=(0, 5))
        info_frame.pack(fill=tk.X)
        ttk.Label(info_frame, textvariable=self.pdf_info_var, font=("Segoe UI", 9, "italic")).pack(side=tk.LEFT)

        # Área de trabajo por pestañas (resultado + visualizador integrado)
        self.workspace_notebook = ttk.Notebook(content_area)
        self.workspace_notebook.pack(fill=tk.BOTH, expand=True)

        output_tab = ttk.Frame(self.workspace_notebook, padding=10)
        self.workspace_notebook.add(output_tab, text="📜 Resultado / Texto")
        self.output_text = scrolledtext.ScrolledText(
            output_tab,
            height=10,
            font=("Consolas", 10),
            borderwidth=1,
            relief=tk.FLAT,
            bg="#ffffff",
        )
        self.output_text.pack(fill=tk.BOTH, expand=True)

        preview_tab = ttk.Frame(self.workspace_notebook, padding=10)
        self.workspace_notebook.add(preview_tab, text="👁️ Vista")
        self.preview_tab_index = 1
        preview_tab.rowconfigure(1, weight=1)
        preview_tab.columnconfigure(0, weight=1)
        self.workspace_notebook.bind("<ButtonRelease-1>", self._on_notebook_click, add="+")

        preview_toolbar = ttk.Frame(preview_tab)
        preview_toolbar.grid(row=0, column=0, sticky="ew", pady=(0, 8))
        self.preview_title_var = tk.StringVar(value="Selecciona un archivo de la lista para previsualizar.")
        ttk.Label(preview_toolbar, textvariable=self.preview_title_var).pack(side=tk.LEFT)
        self.edit_preview_btn = create_button(preview_toolbar, "✍️ Editar este PDF", self._edit_previewed_pdf, style="Accent.TButton")
        self.edit_preview_btn.pack(side=tk.RIGHT)

        viewer_container = ttk.Frame(preview_tab)
        viewer_container.grid(row=1, column=0, sticky="nsew")
        viewer_container.rowconfigure(0, weight=1)
        viewer_container.columnconfigure(0, weight=1)

        self.preview_canvas = tk.Canvas(viewer_container, bg="#f4f4f4", highlightthickness=0)
        self.preview_canvas.grid(row=0, column=0, sticky="nsew")
        self.preview_scrollbar = ttk.Scrollbar(viewer_container, orient=tk.VERTICAL, command=self.preview_canvas.yview)
        self.preview_scrollbar.grid(row=0, column=1, sticky="ns")
        self.preview_canvas.configure(yscrollcommand=self.preview_scrollbar.set)
        self.preview_canvas.bind("<MouseWheel>", self._on_preview_mousewheel)
        self.preview_canvas.bind("<Button-4>", self._on_preview_mousewheel)
        self.preview_canvas.bind("<Button-5>", self._on_preview_mousewheel)

        self.preview_word_text = tk.Text(
            viewer_container,
            wrap=tk.WORD,
            font=("Calibri", 11),
            bg="#ffffff",
            relief=tk.FLAT,
            borderwidth=0,
        )
        self.preview_word_text.tag_configure("title", font=("Segoe UI", 12, "bold"), foreground="#1a73e8")
        self.preview_word_text.tag_configure("heading1", font=("Segoe UI", 14, "bold"))
        self.preview_word_text.tag_configure("heading2", font=("Segoe UI", 12, "bold"))
        self.preview_word_text.tag_configure("heading3", font=("Segoe UI", 11, "bold"))
        self.preview_word_text.tag_configure("bold", font=("Calibri", 11, "bold"))
        self.preview_word_text.tag_configure("italic", font=("Calibri", 11, "italic"))
        self.preview_word_text.tag_configure("table", font=("Consolas", 10), foreground="#3c4043")
        self.preview_word_text.configure(state=tk.DISABLED)

        # 3. Barra de estado
        status_bar = ttk.Frame(self, relief=tk.SUNKEN, padding=(10, 2))
        status_bar.pack(fill=tk.X, side=tk.BOTTOM)
        ttk.Label(status_bar, textvariable=self.status_var, font=("Segoe UI", 8)).pack(side=tk.LEFT)

    def _set_status(self, message: str) -> None:
        """Actualiza el mensaje de la barra de estado."""
        self.status_var.set(f"● {message}")

    def _refresh_file_list(self) -> None:
        """Actualiza la lista visual de archivos cargados."""
        self.files_list.delete(0, tk.END)
        for index, file_path in enumerate(self.loaded_files, start=1):
            self.files_list.insert(tk.END, f" {index:02d}. {file_path.name}")

    def _selected_pdf_index(self) -> int | None:
        """Obtiene el índice del PDF seleccionado en la lista."""
        selected = self.files_list.curselection()
        if not selected:
            return None
        return selected[0]

    def _on_file_selection(self, _event=None) -> None:
        """Evento al seleccionar un archivo de la lista."""
        selected_idx = self._selected_pdf_index()
        if selected_idx is None:
            self.pdf_info_var.set("Selecciona un PDF de la lista")
            return

        input_file = self.loaded_files[selected_idx]
        if input_file.suffix.lower() == ".docx":
            self.pdf_info_var.set(f"📌 Word: {input_file.name}")
            self._preview_selected_file(input_file)
            return

        try:
            total_pages = self.pdf_service.get_total_pages(input_file)
            self.pdf_info_var.set(f"📌 PDF: {input_file.name} | 📖 Páginas: {total_pages}")
        except Exception:
            # Manejar PdfStreamError u otros errores de lectura de metadatos de forma silenciosa
            self.pdf_info_var.set(f"📌 PDF: {input_file.name} | 📖 (Info no disponible)")
        finally:
            self._preview_selected_file(input_file)

    def _move_selected_up(self) -> None:
        """Sube un nivel el archivo seleccionado en la lista."""
        idx = self._selected_pdf_index()
        if idx is None:
            messagebox.showwarning("Orden", "Selecciona un PDF para mover.")
            return
        if idx == 0:
            return

        self.loaded_files[idx - 1], self.loaded_files[idx] = self.loaded_files[idx], self.loaded_files[idx - 1]
        self._refresh_file_list()
        self.files_list.selection_set(idx - 1)
        self._on_file_selection()
        self._set_status("Orden de unión actualizado")

    def _move_selected_down(self) -> None:
        """Baja un nivel el archivo seleccionado en la lista."""
        idx = self._selected_pdf_index()
        if idx is None:
            messagebox.showwarning("Orden", "Selecciona un archivo para mover.")
            return
        if idx == len(self.loaded_files) - 1:
            return

        self.loaded_files[idx + 1], self.loaded_files[idx] = self.loaded_files[idx], self.loaded_files[idx + 1]
        self._refresh_file_list()
        self.files_list.selection_set(idx + 1)
        self._on_file_selection()
        self._set_status("Orden actualizado")

    def _remove_selected_file(self) -> None:
        """Elimina el archivo seleccionado de la lista."""
        idx = self._selected_pdf_index()
        if idx is None:
            messagebox.showwarning("Quitar", "Selecciona un archivo para quitar.")
            return
        
        removed_name = self.loaded_files[idx].name
        del self.loaded_files[idx]
        self._refresh_file_list()
        
        if self.loaded_files:
            new_idx = min(idx, len(self.loaded_files) - 1)
            self.files_list.selection_set(new_idx)
        
        self._on_file_selection()
        self._set_status(f"Archivo '{removed_name}' quitado")

    def _clear_file_list(self) -> None:
        """Limpia todos los archivos de la lista."""
        if not self.loaded_files:
            return
            
        if messagebox.askyesno("Limpiar", "¿Estás seguro de que quieres quitar todos los archivos?"):
            self.loaded_files = []
            self._refresh_file_list()
            self._on_file_selection()
            self._set_status("Lista de archivos vaciada")

    def _load_files(self) -> None:
        """Abre el diálogo para cargar uno o más archivos PDF o Word."""
        selected = filedialog.askopenfilenames(
            title="Selecciona archivos PDF o Word",
            filetypes=[("Archivos permitidos", "*.pdf *.docx"), ("PDF files", "*.pdf"), ("Word files", "*.docx")],
        )
        if not selected:
            return

        try:
            files = self.file_service.validate_mixed_inputs(list(selected))
            # Añadimos a los archivos ya cargados en lugar de reemplazarlos
            self.loaded_files.extend(files)
            self._refresh_file_list()
            if self.loaded_files:
                self.files_list.selection_set(len(self.loaded_files) - len(files))
                self._on_file_selection()
            self._set_status(f"{len(files)} archivo(s) añadido(s) correctamente")
        except Exception as exc:
            messagebox.showerror("Error", human_error(exc))
            self._set_status("Error al cargar archivos")

    def _merge_pdfs(self) -> None:
        """Une los PDFs de la lista (ignora archivos Word)."""
        pdf_files = [f for f in self.loaded_files if f.suffix.lower() == ".pdf"]
        if len(pdf_files) < 2:
            messagebox.showwarning("Unir", "Carga al menos 2 archivos PDF para unir.")
            return
        
        target = filedialog.asksaveasfilename(
            title="Guardar PDF unido",
            defaultextension=".pdf",
            filetypes=[("PDF files", "*.pdf")],
        )
        if not target:
            return

        try:
            output = self.file_service.prepare_output_path(str(ensure_pdf_extension(target)))
            self.pdf_service.merge_pdfs(pdf_files, output)
            self._set_status(f"PDF unido guardado exitosamente")
            messagebox.showinfo("Éxito", f"PDF unido generado en:\n{output.name}")
        except Exception as exc:
            messagebox.showerror("Error", human_error(exc))
            self._set_status("Error al unir documentos")

    def _split_pdf(self) -> None:
        """Abre las opciones de división (ignora archivos Word)."""
        if not self.loaded_files:
            messagebox.showwarning("Dividir", "No hay archivos cargados.")
            return

        selected_idx = self._selected_pdf_index()
        
        # Filtrar solo los PDFs para el contexto y proceso
        pdf_files = [f for f in self.loaded_files if f.suffix.lower() == ".pdf"]
        
        if not pdf_files:
            messagebox.showwarning("Dividir", "No hay archivos PDF cargados para esta acción.")
            return

        # Contexto para el diálogo
        if selected_idx is not None:
            input_pdf_context = self.loaded_files[selected_idx]
            if input_pdf_context.suffix.lower() != ".pdf":
                input_pdf_context = pdf_files[0]
        else:
            input_pdf_context = pdf_files[0]

        split_options = self._ask_split_options(input_pdf_context)
        if split_options is None:
            return

        all_files = split_options.get("all_files", False)
        if not all_files and (selected_idx is None or self.loaded_files[selected_idx].suffix.lower() != ".pdf"):
            messagebox.showwarning("Dividir", "Selecciona un archivo PDF de la lista o marca 'Aplicar a todos'.")
            return

        output_dir = filedialog.askdirectory(title="Carpeta de salida")
        if not output_dir:
            return

        try:
            mode = split_options["mode"]
            files_to_process = pdf_files if all_files else [self.loaded_files[selected_idx]]
            total_created = 0

            for input_pdf in files_to_process:
                if mode == "range":
                    created = self.pdf_service.split_pdf(
                        input_path=input_pdf,
                        output_dir=Path(output_dir),
                        start_page=split_options["start_page"],
                        end_page=split_options["end_page"],
                    )
                elif mode == "parts":
                    created = self.pdf_service.split_pdf_by_parts(
                        input_path=input_pdf,
                        output_dir=Path(output_dir),
                        num_parts=split_options["num_parts"],
                    )
                else:
                    created = self.pdf_service.extract_page_ranges(
                        input_path=input_pdf,
                        output_dir=Path(output_dir),
                        ranges_input=str(split_options["ranges_input"]),
                    )
                total_created += len(created)

            self._set_status(f"Se generaron {total_created} archivo(s)")
            messagebox.showinfo("Éxito", f"División completada: {total_created} archivos creados.")
        except Exception as exc:
            messagebox.showerror("Error", human_error(exc))
            self._set_status("Error en la división")

    def _extract_content(self) -> None:
        """Abre el diálogo de extracción de contenido o conversión a Word."""
        if not self.loaded_files:
            messagebox.showwarning("Extraer", "No hay archivos cargados.")
            return

        selected_idx = self._selected_pdf_index()
        options = self._ask_content_mode()
        if not options:
            return
        
        mode = options["mode"]
        all_files = options["all_files"]

        if not all_files and selected_idx is None:
            messagebox.showwarning("Extraer", "Selecciona un PDF o activa 'Aplicar a todos'.")
            return

        try:
            if all_files:
                self._process_batch_extraction(mode)
            else:
                self._process_single_extraction(self.loaded_files[selected_idx], mode)
        except RuntimeError as exc:
            messagebox.showwarning("Dependencia", str(exc))
            self._set_status("Falta Pillow para imágenes")
        except Exception as exc:
            messagebox.showerror("Error", human_error(exc))
            self._set_status("Error en el proceso")

    def _process_batch_extraction(self, mode: str) -> None:
        """Procesa todos los PDFs cargados (ignora archivos Word)."""
        pdf_files = [f for f in self.loaded_files if f.suffix.lower() == ".pdf"]
        if not pdf_files:
            messagebox.showwarning("Extraer", "No hay archivos PDF cargados.")
            return

        if mode == "text":
            blocks: list[str] = []
            for pdf in pdf_files:
                text = self.pdf_service.extract_text(pdf)
                blocks.append(f"📄 Documento: {pdf.name}\n{text or '[Sin texto]'}")
            combined = "\n\n" + "="*40 + "\n\n".join(blocks).strip()
            self.output_text.delete("1.0", tk.END)
            self.output_text.insert(tk.END, combined)
            
            output_file = filedialog.asksaveasfilename(
                title="Guardar texto combinado",
                defaultextension=".txt",
                initialfile="textos_combinados.txt",
                filetypes=[("Text files", "*.txt")],
            )
            if output_file:
                Path(output_file).write_text(combined, encoding="utf-8")
                self._set_status(f"Texto guardado en {Path(output_file).name}")
        
        elif mode == "text_images":
            base_dir = filedialog.askdirectory(title="Carpeta de destino para batch")
            if not base_dir: return
            for pdf in pdf_files:
                self.pdf_service.extract_text_and_images(pdf, Path(base_dir) / f"{pdf.stem}_extract")
            messagebox.showinfo("Éxito", "Extracción masiva completada.")
            self._set_status("Batch de imágenes y texto finalizado")
            
        elif mode in {"word_quick", "word_advanced"}:
            output_dir = filedialog.askdirectory(title="Carpeta para archivos Word")
            if not output_dir: return
            conv_mode = "advanced" if mode == "word_advanced" else "quick"
            for pdf in pdf_files:
                self.pdf_service.convert_pdf_to_docx(pdf, Path(output_dir) / f"{pdf.stem}.docx", mode=conv_mode)
            messagebox.showinfo("Éxito", f"Se convirtieron {len(pdf_files)} archivos a Word.")
            self._set_status("Conversión masiva a Word completada")

    def _process_single_extraction(self, input_pdf: Path, mode: str) -> None:
        """Procesa un solo archivo según el modo elegido."""
        if mode == "text":
            text = self.pdf_service.extract_text(input_pdf)
            self.output_text.delete("1.0", tk.END)
            self.output_text.insert(tk.END, text or "[Sin texto extraíble]")
            self._set_status(f"Texto extraído de {input_pdf.name}")
            
            output_file = filedialog.asksaveasfilename(
                title="Guardar texto",
                defaultextension=".txt",
                initialfile=f"{input_pdf.stem}.txt",
                filetypes=[("Text files", "*.txt")],
            )
            if output_file:
                Path(output_file).write_text(text, encoding="utf-8")
        
        elif mode == "text_images":
            output_dir = filedialog.askdirectory(title="Carpeta de salida")
            if not output_dir: return
            res = self.pdf_service.extract_text_and_images(input_pdf, Path(output_dir) / f"{input_pdf.stem}_extract")
            messagebox.showinfo("Extracción", f"Finalizado. Imágenes: {res['image_count']}")
            self._set_status("Extracción completada")

        elif mode in {"word_quick", "word_advanced"}:
            output_dir = filedialog.askdirectory(title="Carpeta para Word")
            if not output_dir: return
            conv_mode = "advanced" if mode == "word_advanced" else "quick"
            out_path = self.pdf_service.convert_pdf_to_docx(input_pdf, Path(output_dir) / f"{input_pdf.stem}.docx", mode=conv_mode)
            messagebox.showinfo("Word", f"Archivo generado:\n{out_path.name}")
            self._set_status("Conversión a Word completada")

    def _word_to_pdf(self) -> None:
        """Convierte archivos Word (.docx) cargados a PDF."""
        if not self.loaded_files:
            messagebox.showwarning("Word a PDF", "No hay archivos cargados.")
            return

        selected_idx = self._selected_pdf_index()
        word_files = [f for f in self.loaded_files if f.suffix.lower() == ".docx"]
        
        if not word_files:
            messagebox.showwarning("Word a PDF", "No hay archivos Word (.docx) cargados.")
            return

        # Diálogo para elegir modo (seleccionado o todos)
        mode_options = self._ask_word_to_pdf_mode()
        if not mode_options:
            return
            
        all_files = mode_options.get("all_files", False)
        if not all_files and (selected_idx is None or self.loaded_files[selected_idx].suffix.lower() != ".docx"):
            messagebox.showwarning("Word a PDF", "Selecciona un archivo Word de la lista o marca 'Aplicar a todos'.")
            return

        output_dir = filedialog.askdirectory(title="Selecciona carpeta de salida para los PDFs")
        if not output_dir:
            return

        try:
            files_to_process = word_files if all_files else [self.loaded_files[selected_idx]]
            self._set_status(f"Convirtiendo {len(files_to_process)} archivo(s) Word a PDF...")
            self.master.update_idletasks()

            generated = []
            for docx_path in files_to_process:
                out_path = Path(output_dir) / f"{docx_path.stem}.pdf"
                res = self.pdf_service.convert_docx_to_pdf(docx_path, out_path)
                generated.append(res)
            
            self._set_status(f"Se convirtieron {len(generated)} archivos a PDF")
            messagebox.showinfo("Éxito", f"Se han generado {len(generated)} archivo(s) PDF en:\n{output_dir}")
            
        except Exception as exc:
            messagebox.showerror("Error", human_error(exc))
            self._set_status("Error en la conversión Word a PDF")

    def _ask_word_to_pdf_mode(self) -> dict[str, bool] | None:
        """Diálogo simple para elegir entre convertir uno o todos los archivos Word."""
        dialog = tk.Toplevel(self)
        dialog.title("📘 Configurar Word a PDF")
        dialog.grab_set()
        dialog.resizable(False, False)

        all_files_var = tk.BooleanVar(value=False)
        result: dict[str, bool] = {}

        main_frame = ttk.Frame(dialog, padding=20)
        main_frame.pack()

        ttk.Label(main_frame, text="Conversión de Word a PDF", font=("Segoe UI", 10, "bold")).pack(anchor=tk.W, pady=(0, 10))
        ttk.Label(main_frame, text="Esta acción solo procesará los archivos .docx cargados.", font=("Segoe UI", 8, "italic")).pack(anchor=tk.W, pady=(0, 10))
        
        ttk.Checkbutton(main_frame, text="Aplicar a todos los archivos Word cargados", variable=all_files_var).pack(anchor=tk.W, pady=10)

        btn_frame = ttk.Frame(main_frame, padding=(0, 15, 0, 0))
        btn_frame.pack(fill=tk.X)

        def confirm() -> None:
            result["all_files"] = all_files_var.get()
            dialog.destroy()

        create_button(btn_frame, "✅ Aceptar", confirm, style="Primary.TButton").pack(side=tk.RIGHT, padx=5)
        create_button(btn_frame, "❌ Cancelar", dialog.destroy).pack(side=tk.RIGHT)

        dialog.wait_window()
        return result if result else None

    def _ask_content_mode(self) -> dict[str, str | bool] | None:
        """Muestra el diálogo para elegir el modo de extracción/conversión."""
        dialog = tk.Toplevel(self)
        dialog.title("Configurar Extracción")
        dialog.grab_set()
        dialog.resizable(False, False)

        mode_var = tk.StringVar(value="text")
        all_files_var = tk.BooleanVar(value=False)
        result: dict[str, str | bool] = {}

        main_frame = ttk.Frame(dialog, padding=20)
        main_frame.pack()

        ttk.Label(main_frame, text="¿Qué deseas hacer?", font=("Segoe UI", 10, "bold")).pack(anchor=tk.W, pady=(0, 10))
        
        options = [
            ("Solo Texto (.txt)", "text"),
            ("Texto + Imágenes", "text_images"),
            ("Word (Rápido)", "word_quick"),
            ("Word (Avanzado - Recomendado)", "word_advanced")
        ]

        for text, value in options:
            ttk.Radiobutton(main_frame, text=text, variable=mode_var, value=value).pack(anchor=tk.W, pady=2)

        ttk.Separator(main_frame, orient=tk.HORIZONTAL).pack(fill=tk.X, pady=15)
        
        ttk.Checkbutton(main_frame, text="Aplicar a todos los archivos cargados", variable=all_files_var).pack(anchor=tk.W)

        btn_frame = ttk.Frame(main_frame, padding=(0, 15, 0, 0))
        btn_frame.pack(fill=tk.X)

        def confirm() -> None:
            result["mode"] = mode_var.get()
            result["all_files"] = all_files_var.get()
            dialog.destroy()

        create_button(btn_frame, "✅ Aceptar", confirm, style="Primary.TButton").pack(side=tk.RIGHT, padx=5)
        create_button(btn_frame, "❌ Cancelar", dialog.destroy).pack(side=tk.RIGHT)

        dialog.wait_window()
        return result if result else None

    def _preview_selected_file(self, input_file: Path) -> None:
        """Visualiza automáticamente el archivo seleccionado en la pestaña integrada."""
        self.workspace_notebook.select(1)
        try:
            if input_file.suffix.lower() == ".pdf":
                self._render_pdf_in_app_viewer(input_file)
                self._set_status(f"Visualizando {input_file.name} en pestaña integrada")
            else:
                self._render_word_in_app_viewer(input_file)
                self._set_status(f"Visualizando {input_file.name} en pestaña integrada")
        except Exception as exc:
            messagebox.showerror("Error", human_error(exc))
            self._set_status("Error al visualizar archivo")

    def _render_pdf_in_app_viewer(self, input_pdf: Path, *, display_name: str | None = None, allow_edit: bool = True) -> None:
        visible_name = display_name or input_pdf.name
        self._set_preview_tab_title(visible_name)
        self.preview_canvas.delete("all")
        self.preview_images = []
        self.preview_current_pdf = input_pdf if allow_edit else None
        self.preview_title_var.set(f"Vista integrada: {visible_name}")
        self._show_pdf_preview_mode(allow_edit=allow_edit)

        total_pages = self.pdf_service.get_total_pages(input_pdf)
        y_offset = 12
        page_gap = 16
        max_width = 0
        for page_number in range(1, total_pages + 1):
            image_bytes = self.pdf_service.render_pdf_page_preview(input_pdf, page_number=page_number, zoom=1.1)
            encoded = base64.b64encode(image_bytes).decode("ascii")
            photo = tk.PhotoImage(data=encoded)
            self.preview_images.append(photo)

            self.preview_canvas.create_text(
                12,
                y_offset,
                anchor=tk.NW,
                text=f"Página {page_number}/{total_pages}",
                font=("Segoe UI", 9, "bold"),
                fill="#202124",
            )
            y_offset += 20
            self.preview_canvas.create_image(12, y_offset, anchor=tk.NW, image=photo)
            y_offset += photo.height() + page_gap
            max_width = max(max_width, photo.width())

        self.preview_canvas.configure(scrollregion=(0, 0, max_width + 32, y_offset))
        self.preview_canvas.yview_moveto(0)

    def _render_word_in_app_viewer(self, input_docx: Path) -> None:
        if self._try_render_word_with_spire(input_docx):
            return

        self._set_preview_tab_title(input_docx.name)
        self.preview_images = []
        self.preview_current_pdf = None
        self.preview_title_var.set(f"Vista integrada: {input_docx.name} (Word / fallback python-docx)")
        self._show_word_preview_mode()
        self.preview_word_text.configure(state=tk.NORMAL)
        self.preview_word_text.delete("1.0", tk.END)

        doc = Document(str(input_docx))
        self.preview_word_text.insert(tk.END, f"{input_docx.name}\n", "title")
        self.preview_word_text.insert(tk.END, "-" * max(40, len(input_docx.name)) + "\n\n")

        for paragraph in doc.paragraphs:
            style_name = (paragraph.style.name or "").lower() if paragraph.style else ""
            style_tag = ()
            prefix = ""
            if style_name.startswith("heading 1"):
                style_tag = ("heading1",)
            elif style_name.startswith("heading 2"):
                style_tag = ("heading2",)
            elif style_name.startswith("heading 3"):
                style_tag = ("heading3",)
            elif "list" in style_name or "bullet" in style_name:
                prefix = "• "

            if prefix:
                self.preview_word_text.insert(tk.END, prefix)

            if not paragraph.runs:
                plain_text = paragraph.text.strip()
                if plain_text:
                    self.preview_word_text.insert(tk.END, f"{plain_text}\n", style_tag)
                continue

            for run in paragraph.runs:
                run_tags: tuple[str, ...] = style_tag
                if run.bold:
                    run_tags += ("bold",)
                if run.italic:
                    run_tags += ("italic",)
                self.preview_word_text.insert(tk.END, run.text, run_tags)
            self.preview_word_text.insert(tk.END, "\n")

        for table in doc.tables:
            self.preview_word_text.insert(tk.END, "\n[Tabla]\n", ("table",))
            for row in table.rows:
                row_values = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                line = " | ".join(value if value else "-" for value in row_values)
                self.preview_word_text.insert(tk.END, f"{line}\n", ("table",))

        # Inserta imágenes embebidas del documento si existen (sin conversión a PDF)
        self.preview_word_text.insert(tk.END, "\n")
        for rel in doc.part.rels.values():
            if "image" not in rel.reltype:
                continue
            image_blob = rel.target_part.blob
            photo = self._photo_from_image_blob(image_blob)
            if photo is None:
                continue
            self.preview_images.append(photo)
            self.preview_word_text.insert(tk.END, "[Imagen]\n", ("table",))
            self.preview_word_text.image_create(tk.END, image=photo)
            self.preview_word_text.insert(tk.END, "\n\n")

        self.preview_word_text.configure(state=tk.DISABLED)
        self.preview_word_text.yview_moveto(0)

    def _try_render_word_with_spire(self, input_docx: Path) -> bool:
        """
        Renderiza DOCX con Spire.Doc para vista previa visual.
        Devuelve False si Spire.Doc no está disponible o falla el render.
        """
        try:
            from spire.doc import Document as SpireDocument  # type: ignore
            from spire.doc import FileFormat  # type: ignore
        except Exception:
            if not self._spire_notice_shown:
                self._set_status("Spire.Doc no disponible: usando vista Word fallback (python-docx)")
                self._spire_notice_shown = True
            return False

        self._cleanup_preview_temp_files()
        try:
            temp_dir = Path(tempfile.mkdtemp(prefix="word_preview_spire_"))
            temp_pdf = temp_dir / f"{input_docx.stem}_preview.pdf"
            doc = SpireDocument()
            doc.LoadFromFile(str(input_docx))
            doc.SaveToFile(str(temp_pdf), FileFormat.PDF)
            self.preview_temp_dir = temp_dir
        except Exception:
            self._cleanup_preview_temp_files()
            return False

        self._render_pdf_in_app_viewer(temp_pdf, display_name=input_docx.name, allow_edit=False)
        self.preview_title_var.set(f"Vista integrada: {input_docx.name} (Word / Spire.Doc)")
        self._set_status("Vista Word renderizada con Spire.Doc")
        return True

    def _on_preview_mousewheel(self, event) -> None:
        if getattr(event, "num", None) == 4:
            self.preview_canvas.yview_scroll(-3, "units")
        elif getattr(event, "num", None) == 5:
            self.preview_canvas.yview_scroll(3, "units")
        else:
            delta_units = -1 if event.delta > 0 else 1
            self.preview_canvas.yview_scroll(delta_units * 3, "units")

    def _edit_previewed_pdf(self) -> None:
        if self.preview_current_pdf is None:
            messagebox.showwarning("Editar PDF", "Selecciona y visualiza un PDF para editar.")
            return
        try:
            subprocess.Popen([sys.executable, "-m", "src.ui.pdf_editor_app", str(self.preview_current_pdf)])
            self._set_status(f"Editor PDF abierto para {self.preview_current_pdf.name}")
        except Exception as exc:
            messagebox.showerror("Error", human_error(exc))
            self._set_status("Error al abrir editor PDF")

    def _clear_preview(self) -> None:
        self._cleanup_preview_temp_files()
        self.preview_canvas.delete("all")
        self.preview_images = []
        self.preview_current_pdf = None
        self.preview_word_text.configure(state=tk.NORMAL)
        self.preview_word_text.delete("1.0", tk.END)
        self.preview_word_text.configure(state=tk.DISABLED)
        self.preview_title_var.set("Visualización cerrada. Selecciona un archivo de la lista para visualizar.")
        self._set_preview_tab_title("👁️ Vista")
        self._show_pdf_preview_mode(allow_edit=False)
        self.workspace_notebook.select(0)
        self._set_status("Visualización cerrada")

    def _set_preview_tab_title(self, label: str) -> None:
        trimmed = label if len(label) <= 32 else f"{label[:29]}..."
        self.preview_tab_display_name = trimmed
        self.workspace_notebook.tab(self.preview_tab_index, text=f"{trimmed}  ✕")

    def _show_word_preview_mode(self) -> None:
        self.preview_canvas.grid_remove()
        self.preview_word_text.grid(row=0, column=0, sticky="nsew")
        self.edit_preview_btn.pack_forget()
        self.preview_scrollbar.grid()
        self.preview_scrollbar.configure(command=self.preview_word_text.yview)
        self.preview_word_text.configure(yscrollcommand=self.preview_scrollbar.set)

    def _show_pdf_preview_mode(self, *, allow_edit: bool) -> None:
        self.preview_word_text.grid_remove()
        self.preview_canvas.grid(row=0, column=0, sticky="nsew")
        self.preview_scrollbar.grid()
        if allow_edit and not self.edit_preview_btn.winfo_ismapped():
            self.edit_preview_btn.pack(side=tk.RIGHT)
        if not allow_edit and self.edit_preview_btn.winfo_ismapped():
            self.edit_preview_btn.pack_forget()
        self.preview_scrollbar.configure(command=self.preview_canvas.yview)
        self.preview_canvas.configure(yscrollcommand=self.preview_scrollbar.set)

    def _photo_from_image_blob(self, image_blob: bytes) -> tk.PhotoImage | None:
        try:
            encoded = base64.b64encode(image_blob).decode("ascii")
            return tk.PhotoImage(data=encoded)
        except Exception:
            return None

    def _on_notebook_click(self, event) -> None:
        try:
            tab_id = self.workspace_notebook.index(f"@{event.x},{event.y}")
        except tk.TclError:
            return
        if tab_id != self.preview_tab_index:
            return
        try:
            x1, _y1, width, _height = self.workspace_notebook.bbox(tab_id)
        except tk.TclError:
            return
        close_text = "✕"
        close_width = 22 if close_text else 22
        if event.x >= (x1 + width - close_width):
            self._clear_preview()

    def _cleanup_preview_temp_files(self) -> None:
        if self.preview_temp_dir is not None:
            shutil.rmtree(self.preview_temp_dir, ignore_errors=True)
            self.preview_temp_dir = None

    def _edit_pdf(self) -> None:
        """Abre el editor de PDF avanzado (PyQt6) para edición por superposiciones."""
        if not self.loaded_files:
            messagebox.showwarning("Editar PDF", "No hay archivos cargados.")
            return

        selected_idx = self._selected_pdf_index()
        if selected_idx is None:
            messagebox.showwarning("Editar PDF", "Selecciona un PDF de la lista.")
            return

        input_pdf = self.loaded_files[selected_idx]
        if input_pdf.suffix.lower() != ".pdf":
            messagebox.showwarning("Editar PDF", "La edición solo está disponible para archivos PDF.")
            return

        try:
            subprocess.Popen([sys.executable, "-m", "src.ui.pdf_editor_app", str(input_pdf)])
            self._set_status(f"Editor PDF abierto para {input_pdf.name}")
        except Exception as exc:
            messagebox.showerror("Error", human_error(exc))
            self._set_status("Error al abrir editor PDF")

    def _open_pdf_edit_dialog(self, input_pdf: Path) -> None:
        """Diálogo gráfico para visualizar y editar el PDF seleccionado."""
        dialog = tk.Toplevel(self)
        dialog.title(f"✍️ Editor PDF - {input_pdf.name}")
        dialog.geometry("1100x760")
        dialog.grab_set()

        total_pages = self.pdf_service.get_total_pages(input_pdf)
        temp_dir = Path(tempfile.mkdtemp(prefix="pdf_editor_"))
        working_pdf = temp_dir / "working.pdf"
        shutil.copy2(input_pdf, working_pdf)

        mode_var = tk.StringVar(value="replace")
        page_var = tk.IntVar(value=1)
        signature_image_var = tk.StringVar(value="")
        clicked_pdf_coords: dict[str, float] = {"x": 72.0, "y": 72.0}
        render_meta: dict[str, float] = {"img_w": 1.0, "img_h": 1.0, "pdf_w": 1.0, "pdf_h": 1.0}

        main = ttk.Frame(dialog, padding=10)
        main.pack(fill=tk.BOTH, expand=True)
        main.columnconfigure(0, weight=3)
        main.columnconfigure(1, weight=2)
        main.rowconfigure(0, weight=1)

        preview_frame = ttk.LabelFrame(main, text="Vista previa del PDF", padding=10)
        preview_frame.grid(row=0, column=0, sticky="nsew", padx=(0, 10))
        preview_frame.rowconfigure(1, weight=1)
        preview_frame.columnconfigure(0, weight=1)

        top_bar = ttk.Frame(preview_frame)
        top_bar.grid(row=0, column=0, sticky="ew", pady=(0, 8))
        ttk.Label(top_bar, text=f"Página (1-{total_pages}):").pack(side=tk.LEFT)
        ttk.Spinbox(top_bar, from_=1, to=total_pages, textvariable=page_var, width=6).pack(
            side=tk.LEFT, padx=(6, 8)
        )
        ttk.Label(top_bar, text="Selecciona modo y haz click en el PDF.").pack(side=tk.LEFT, padx=(8, 0))

        canvas = tk.Canvas(preview_frame, bg="#f4f4f4", highlightthickness=0)
        canvas.grid(row=1, column=0, sticky="nsew")
        page_text = scrolledtext.ScrolledText(preview_frame, height=8, font=("Consolas", 9))
        page_text.grid(row=2, column=0, sticky="ew", pady=(8, 0))

        controls = ttk.LabelFrame(main, text="Configuración de edición", padding=10)
        controls.grid(row=0, column=1, sticky="nsew")
        controls.columnconfigure(1, weight=1)

        ttk.Label(controls, text="Operación:").grid(row=0, column=0, sticky="w", pady=4)
        mode_combo = ttk.Combobox(
            controls,
            textvariable=mode_var,
            values=["replace", "add", "sign", "edit_click", "image"],
            state="readonly",
        )
        mode_combo.grid(row=0, column=1, sticky="ew", pady=4)

        ttk.Label(controls, text="Buscar texto (replace global):").grid(row=1, column=0, sticky="w", pady=4)
        search_entry = ttk.Entry(controls)
        search_entry.grid(row=1, column=1, sticky="ew", pady=4)

        ttk.Label(controls, text="Reemplazar por:").grid(row=2, column=0, sticky="w", pady=4)
        replace_entry = ttk.Entry(controls)
        replace_entry.grid(row=2, column=1, sticky="ew", pady=4)

        ttk.Label(controls, text="Texto para add/edit_click:").grid(row=3, column=0, sticky="w", pady=4)
        add_text_entry = ttk.Entry(controls)
        add_text_entry.grid(row=3, column=1, sticky="ew", pady=4)

        ttk.Label(controls, text="Firmante (modo sign):").grid(row=4, column=0, sticky="w", pady=4)
        signer_entry = ttk.Entry(controls)
        signer_entry.grid(row=4, column=1, sticky="ew", pady=4)

        ttk.Label(controls, text="Imagen (sign/image):").grid(row=5, column=0, sticky="w", pady=4)
        ttk.Entry(controls, textvariable=signature_image_var).grid(row=5, column=1, sticky="ew", pady=4)
        click_info = ttk.Label(controls, text="Posición click PDF: x=72.0, y=72.0")
        click_info.grid(row=6, column=0, columnspan=2, sticky="w", pady=4)

        def select_signature_image() -> None:
            chosen = filedialog.askopenfilename(
                title="Selecciona imagen de firma",
                filetypes=[("Imágenes", "*.png *.jpg *.jpeg *.bmp")],
            )
            if chosen:
                signature_image_var.set(chosen)

        create_button(controls, "Seleccionar imagen", select_signature_image).grid(
            row=7, column=1, sticky="e", pady=(4, 8)
        )

        def refresh_preview() -> None:
            current_page = page_var.get()
            image_bytes = self.pdf_service.render_pdf_page_preview(working_pdf, page_number=current_page, zoom=1.0)
            encoded = base64.b64encode(image_bytes).decode("ascii")
            photo = tk.PhotoImage(data=encoded)
            canvas.delete("all")
            canvas.create_image(0, 0, anchor=tk.NW, image=photo)
            canvas.image = photo
            canvas.configure(scrollregion=(0, 0, photo.width(), photo.height()))

            pdf_w, pdf_h = self.pdf_service.get_page_size(working_pdf, page_number=current_page)
            render_meta["img_w"] = float(photo.width())
            render_meta["img_h"] = float(photo.height())
            render_meta["pdf_w"] = pdf_w
            render_meta["pdf_h"] = pdf_h

            extracted = self.pdf_service.extract_text_from_page(working_pdf, page_number=current_page) or "[Sin texto]"
            page_text.delete("1.0", tk.END)
            page_text.insert(tk.END, extracted)

        def apply_click_action() -> None:
            temp_out = temp_dir / "working_next.pdf"
            try:
                mode = mode_var.get()
                if mode == "replace":
                    self.pdf_service.replace_text_in_pdf(
                        input_path=working_pdf,
                        output_path=temp_out,
                        search_text=search_entry.get().strip(),
                        replace_text=replace_entry.get().strip(),
                    )
                elif mode == "add":
                    self.pdf_service.add_text_to_pdf(
                        input_path=working_pdf,
                        output_path=temp_out,
                        text=add_text_entry.get().strip(),
                        page_number=page_var.get(),
                        x=clicked_pdf_coords["x"],
                        y=clicked_pdf_coords["y"],
                    )
                elif mode == "sign":
                    signature_image = signature_image_var.get().strip()
                    self.pdf_service.sign_pdf_at_position(
                        input_path=working_pdf,
                        output_path=temp_out,
                        signer_name=signer_entry.get().strip(),
                        page_number=page_var.get(),
                        x=clicked_pdf_coords["x"],
                        y=clicked_pdf_coords["y"],
                        signature_image_path=Path(signature_image) if signature_image else None,
                    )
                elif mode == "image":
                    image_path = signature_image_var.get().strip()
                    self.pdf_service.insert_image_to_pdf(
                        input_path=working_pdf,
                        output_path=temp_out,
                        image_path=Path(image_path),
                        page_number=page_var.get(),
                        x=clicked_pdf_coords["x"],
                        y=clicked_pdf_coords["y"],
                    )
                else:
                    self.pdf_service.replace_text_at_position(
                        input_path=working_pdf,
                        output_path=temp_out,
                        page_number=page_var.get(),
                        x=clicked_pdf_coords["x"],
                        y=clicked_pdf_coords["y"],
                        replacement_text=add_text_entry.get().strip(),
                    )

                temp_out.replace(working_pdf)
                refresh_preview()
                self._set_status("Edición aplicada en vista previa")
            except Exception as exc:
                messagebox.showerror("Error", human_error(exc))
                self._set_status("Error al aplicar edición")

        def save_edited_pdf() -> None:
            output_path_raw = filedialog.asksaveasfilename(
                title="Guardar PDF editado",
                defaultextension=".pdf",
                initialfile=f"{input_pdf.stem}_editado.pdf",
                filetypes=[("PDF files", "*.pdf")],
            )
            if not output_path_raw:
                return
            output_path = self.file_service.prepare_output_path(str(ensure_pdf_extension(output_path_raw)))
            shutil.copy2(working_pdf, output_path)
            self._set_status(f"PDF editado: {output_path.name}")
            messagebox.showinfo("Éxito", f"Archivo generado:\n{output_path}")

        def on_canvas_click(event) -> None:
            if render_meta["img_w"] <= 0 or render_meta["img_h"] <= 0:
                return
            x_img = min(max(event.x, 0), int(render_meta["img_w"]))
            y_img = min(max(event.y, 0), int(render_meta["img_h"]))
            pdf_x = (x_img / render_meta["img_w"]) * render_meta["pdf_w"]
            pdf_y = (y_img / render_meta["img_h"]) * render_meta["pdf_h"]
            clicked_pdf_coords["x"] = pdf_x
            clicked_pdf_coords["y"] = pdf_y
            click_info.configure(text=f"Posición click PDF: x={pdf_x:.1f}, y={pdf_y:.1f}")
            apply_click_action()

        canvas.bind("<Button-1>", on_canvas_click)

        action_row = ttk.Frame(controls)
        action_row.grid(row=8, column=0, columnspan=2, sticky="ew", pady=(12, 0))
        create_button(action_row, "Actualizar vista", refresh_preview).pack(side=tk.LEFT, padx=(0, 6))
        create_button(action_row, "Guardar como", save_edited_pdf, style="Primary.TButton").pack(side=tk.LEFT)

        refresh_preview()

        def on_close() -> None:
            shutil.rmtree(temp_dir, ignore_errors=True)
            dialog.destroy()

        dialog.protocol("WM_DELETE_WINDOW", on_close)

    def _ask_split_options(self, input_pdf: Path) -> dict[str, int | str | None] | None:
        """Muestra el diálogo para configurar la división de un PDF con campos dinámicos."""
        total_pages = self.pdf_service.get_total_pages(input_pdf)
        dialog = tk.Toplevel(self)
        dialog.title("✂️ Configurar División")
        dialog.grab_set()
        dialog.resizable(False, False)

        mode_var = tk.StringVar(value="range")
        all_files_var = tk.BooleanVar(value=False)
        result: dict[str, int | str | bool | None] = {}

        main_frame = ttk.Frame(dialog, padding=20)
        main_frame.pack()

        ttk.Label(main_frame, text=f"Documento: {input_pdf.name}", font=("Segoe UI", 9, "bold")).pack(anchor=tk.W, pady=(0, 10))
        
        # Función para habilitar/deshabilitar campos según el modo
        def update_entries():
            mode = mode_var.get()
            # Modo Rango
            start_ent.configure(state="normal" if mode == "range" else "disabled")
            end_ent.configure(state="normal" if mode == "range" else "disabled")
            # Modo Partes
            parts_ent.configure(state="normal" if mode == "parts" else "disabled")
            # Modo Rangos Múltiples
            multi_ent.configure(state="normal" if mode == "extract_ranges" else "disabled")

        # Modo Rango
        range_frame = ttk.LabelFrame(main_frame, text="1. Por Rango Único", padding=10)
        range_frame.pack(fill=tk.X, pady=5)
        ttk.Radiobutton(range_frame, text="Activar", variable=mode_var, value="range", command=update_entries).pack(anchor=tk.W)
        ttk.Label(range_frame, text="Genera un archivo por cada una de las páginas entre el Inicio y el Fin.", font=("Segoe UI", 8, "italic")).pack(anchor=tk.W)
        
        f1 = ttk.Frame(range_frame)
        f1.pack(fill=tk.X, pady=5)
        ttk.Label(f1, text="Inicio (vacío = 1):").pack(side=tk.LEFT)
        start_ent = ttk.Entry(f1, width=5)
        start_ent.pack(side=tk.LEFT, padx=5)
        ttk.Label(f1, text="Fin (vacío = última):").pack(side=tk.LEFT)
        end_ent = ttk.Entry(f1, width=5)
        end_ent.pack(side=tk.LEFT, padx=5)

        # Modo Partes
        parts_frame = ttk.LabelFrame(main_frame, text="2. Dividir en N Partes", padding=10)
        parts_frame.pack(fill=tk.X, pady=5)
        ttk.Radiobutton(parts_frame, text="Activar", variable=mode_var, value="parts", command=update_entries).pack(anchor=tk.W)
        ttk.Label(parts_frame, text="Divide el PDF en varios archivos de tamaño similar.", font=("Segoe UI", 8, "italic")).pack(anchor=tk.W)
        
        f2 = ttk.Frame(parts_frame)
        f2.pack(fill=tk.X, pady=5)
        ttk.Label(f2, text="Nº de archivos:").pack(side=tk.LEFT)
        parts_ent = ttk.Entry(f2, width=5)
        parts_ent.pack(side=tk.LEFT, padx=5)

        # Modo Rangos Múltiples
        multi_frame = ttk.LabelFrame(main_frame, text="3. Extraer Múltiples Rangos", padding=10)
        multi_frame.pack(fill=tk.X, pady=5)
        ttk.Radiobutton(multi_frame, text="Activar", variable=mode_var, value="extract_ranges", command=update_entries).pack(anchor=tk.W)
        ttk.Label(multi_frame, text="Crea un PDF por cada rango especificado.", font=("Segoe UI", 8, "italic")).pack(anchor=tk.W)
        
        multi_ent = ttk.Entry(multi_frame)
        multi_ent.pack(fill=tk.X, pady=5)
        ttk.Label(multi_frame, text="Ejemplo: 1-3, 5, 10-12 (Obligatorio)", font=("Segoe UI", 8, "italic"), foreground="gray").pack(anchor=tk.W)

        ttk.Separator(main_frame, orient=tk.HORIZONTAL).pack(fill=tk.X, pady=10)
        
        ttk.Checkbutton(main_frame, text="Aplicar a todos los archivos cargados", variable=all_files_var).pack(anchor=tk.W, pady=(0, 10))

        # Inicializar estados
        update_entries()

        def confirm() -> None:
            try:
                mode = mode_var.get()
                result["all_files"] = all_files_var.get()
                if mode == "range":
                    s_raw = start_ent.get().strip()
                    e_raw = end_ent.get().strip()
                    s = int(s_raw) if s_raw else 1
                    e = int(e_raw) if e_raw else total_pages
                    if s < 1 or e > total_pages or s > e:
                        raise ValueError("Rango inválido.")
                    result.update({"mode": "range", "start_page": s, "end_page": e})
                elif mode == "parts":
                    p_raw = parts_ent.get().strip()
                    if not p_raw: raise ValueError("Ingresa el número de partes.")
                    p = int(p_raw)
                    if p < 1 or p > total_pages: raise ValueError("Número de partes inválido.")
                    result.update({"mode": "parts", "num_parts": p})
                else:
                    m_raw = multi_ent.get().strip()
                    if not m_raw: raise ValueError("Ingresa al menos un rango.")
                    result.update({"mode": "extract_ranges", "ranges_input": m_raw})
                dialog.destroy()
            except ValueError as exc:
                messagebox.showwarning("Valor Inválido", str(exc) if str(exc) else "Revisa los campos.")

        btn_f = ttk.Frame(main_frame, padding=(0, 10, 0, 0))
        btn_f.pack(fill=tk.X)
        create_button(btn_f, "✅ Aceptar", confirm, style="Primary.TButton").pack(side=tk.RIGHT, padx=5)
        create_button(btn_f, "❌ Cancelar", dialog.destroy).pack(side=tk.RIGHT)

        dialog.wait_window()
        return result or None
